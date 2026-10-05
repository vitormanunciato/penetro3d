import numpy as np
import openpyxl
import pytest

from penetro3d_core import ErroDeDados, detectar_outliers, ler_falker

LINHA_PROF = 28          # primeira linha de leituras no layout do Falker
COL_PRIMEIRA = 3         # coluna C = medição 1


def test_le_o_exemplo(xlsx_exemplo):
    pontos, prof, CI, data, avisos = ler_falker(xlsx_exemplo)
    assert list(pontos.columns) == ['id', 'lat', 'lon']
    assert len(pontos) == 33 and CI.shape == (21, 33)
    assert prof[0] == 2.5 and prof[-1] == 52.5           # corte comum no perfil interrompido
    assert data == '14/09/2026'
    assert any('52,5' in a for a in avisos)
    assert np.isfinite(CI).all()


def test_pedra_vira_outlier(xlsx_exemplo):
    pontos, prof, CI, _, _ = ler_falker(xlsx_exemplo)
    out = detectar_outliers(pontos, prof, CI)
    assert [o[0] for o in out] == ['8']
    assert out[0][1] > 7000


def test_perfil_muito_curto_e_descartado(xlsx_exemplo, tmp_path):
    wb = openpyxl.load_workbook(xlsx_exemplo)
    ws = wb.active
    for k in range(8, 24):                           # medição 1 só até 20 cm
        ws.cell(row=LINHA_PROF + k, column=COL_PRIMEIRA).value = None
    p = tmp_path / 'curto.xlsx'
    wb.save(p)
    pontos, prof, CI, _, avisos = ler_falker(str(p))
    assert '1' not in set(pontos.id)
    assert len(pontos) == 32
    assert prof[-1] == 52.5                          # um perfil ruim não encurta todos
    assert any('1' in a for a in avisos)


def test_planilha_sem_layout_falker(tmp_path):
    wb = openpyxl.Workbook()
    wb.active['A1'] = 'qualquer coisa'
    p = tmp_path / 'outra.xlsx'
    wb.save(p)
    with pytest.raises(ErroDeDados):
        ler_falker(str(p))
