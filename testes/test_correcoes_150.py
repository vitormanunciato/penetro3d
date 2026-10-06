"""Regressão das falhas encontradas na revisão de outubro/2026 (versão 1.5.0).

Cada teste reproduz um caso que antes dava resultado errado ou quebrava.
"""
import os
import xml.etree.ElementTree as ET
import zipfile

import numpy as np
import openpyxl
import pandas as pd
import pytest
from conftest import EXEMPLOS, anel_txt, escrever_kml

import penetro3d_core as C
from penetro3d_amostragem import _gpx, planejar

LINHA_PROF, COL1 = 28, 3                      # layout do Falker no exemplo
XLSX = os.path.join(EXEMPLOS, 'penetrometria_exemplo.xlsx')
KML = os.path.join(EXEMPLOS, 'talhoes_exemplo.kml')


def _planilha(tmp_path, nome='p.xlsx', **mudar):
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    for (lin, col), v in mudar.get('celulas', {}).items():
        ws.cell(row=lin, column=col).value = v
    if mudar.get('fator'):
        for r in range(LINHA_PROF, LINHA_PROF + 24):
            for c in range(COL1, ws.max_column + 1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, (int, float)):
                    ws.cell(row=r, column=c).value = v * mudar['fator']
    if mudar.get('cabecalho'):
        for c in range(COL1, ws.max_column + 1):
            ws.cell(row=LINHA_PROF - 1, column=c).value = mudar['cabecalho']
    p = tmp_path / nome
    wb.save(p)
    return str(p)


# ───────────────────────────────────────────────────────── leitura da planilha ──
def test_uma_celula_vazia_nao_corta_a_analise(tmp_path):
    p = _planilha(tmp_path, celulas={(LINHA_PROF + 2, COL1 + 4): None})   # 7,5 cm
    pontos, prof, CI, _, avisos = C.ler_falker(p)
    assert prof.max() == 52.5                     # antes: 5 cm para todos
    assert np.isfinite(CI).all()
    assert any('preenchida' in a for a in avisos)


def test_falha_longa_encurta_so_aquele_perfil(tmp_path):
    vazias = {(LINHA_PROF + k, COL1): None for k in range(10, 14)}       # 27,5–35 cm
    p = _planilha(tmp_path, celulas=vazias)
    pontos, prof, CI, _, avisos = C.ler_falker(p)
    assert '1' not in set(pontos.id)              # raso demais: descartado
    assert prof.max() == 52.5
    assert any('falha' in a for a in avisos)


def test_planilha_em_mpa_pelo_cabecalho(tmp_path):
    p = _planilha(tmp_path, fator=0.001, cabecalho='CI (MPa)')
    _, _, CI, _, avisos = C.ler_falker(p)
    ref = C.ler_falker(XLSX)[2]
    assert np.allclose(CI, ref, rtol=1e-6)
    assert any('MPa' in a for a in avisos)


def test_planilha_em_mpa_sem_cabecalho(tmp_path):
    p = _planilha(tmp_path, fator=0.001, cabecalho='CI')
    _, _, CI, _, avisos = C.ler_falker(p)
    assert np.nanmedian(CI) > 500
    assert any('MPa' in a for a in avisos)


def test_virgula_decimal_e_gps_sem_sinal(tmp_path):
    p = _planilha(tmp_path, celulas={(21, COL1): '-25,310700',               # latitude texto
                                     (21, COL1 + 1): 0, (22, COL1 + 1): 0})  # sem GPS
    pontos, *_ , avisos = C.ler_falker(p)
    assert '1' in set(pontos.id)                  # vírgula decimal lida
    assert '2' not in set(pontos.id)              # 0,0 não é coordenada
    assert any('sem coordenada' in a for a in avisos)


def test_coordenada_repetida_avisa(tmp_path):
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    ws.cell(row=21, column=COL1 + 1).value = ws.cell(row=21, column=COL1).value
    ws.cell(row=22, column=COL1 + 1).value = ws.cell(row=22, column=COL1).value
    p = tmp_path / 'dup.xlsx'
    wb.save(p)
    avisos = C.ler_falker(str(p))[4]
    assert any('mesmo lugar' in a and '1 e 2' in a for a in avisos)


# ──────────────────────────────────────────────────────────────── KML / KMZ ──
def test_coordenadas_com_espaco_depois_da_virgula(tmp_path):
    anel = anel_txt(-51.40, -25.30).replace(',', ', ')
    t = C.carregar_talhoes([escrever_kml(tmp_path / 'e.kml',
                            f'<Placemark><name>A</name><Polygon><outerBoundaryIs><LinearRing>'
                            f'<coordinates>{anel}</coordinates></LinearRing></outerBoundaryIs>'
                            f'</Polygon></Placemark>')])
    assert len(t) == 1 and len(t[0]['anel']) == 5


def test_kmz(tmp_path):
    kmz = tmp_path / 'talhoes.kmz'
    with zipfile.ZipFile(kmz, 'w') as z:
        z.write(KML, 'doc.kml')
    assert [t['nome'] for t in C.carregar_talhoes([str(kmz)])] == ['Talhão Norte', 'Talhão Sul']


def test_contorno_desenhado_como_caminho_fechado(tmp_path):
    t = C.carregar_talhoes([escrever_kml(tmp_path / 'l.kml',
                            f'<Placemark><name>Caminho</name><LineString><coordinates>'
                            f'{anel_txt(-51.40, -25.30)}</coordinates></LineString></Placemark>')])
    assert t[0]['nome'] == 'Caminho'


def test_contorno_cruzado_mantem_a_area(tmp_path):
    from pyproj import Transformer
    lon0, lat0, l = -51.40, -25.30, 0.004
    gravata = [(lon0, lat0), (lon0 + l, lat0 + l), (lon0 + l, lat0), (lon0, lat0 + l), (lon0, lat0)]
    tr = Transformer.from_crs('EPSG:4326', 'EPSG:31982', always_xy=True)
    g = C._projetar(gravata, [], tr)
    assert g.area / 1e4 > 8.5                     # antes: 4,46 ha (metade sumia)
    assert C.contorno_corrigido({'anel': gravata, 'ilhas': []})


def test_pastas_nao_colidem(tmp_path):
    corpo = ''.join(f'<Placemark><name>{n}</name><Polygon><outerBoundaryIs><LinearRing>'
                    f'<coordinates>{anel_txt(-51.40 + 0.01 * i, -25.30)}</coordinates>'
                    f'</LinearRing></outerBoundaryIs></Polygon></Placemark>'
                    for i, n in enumerate(['Talhão A', 'talhao a', 'Talhão A.',
                                           'X' * 70, 'X' * 70 + ' outro']))
    t = C.carregar_talhoes([escrever_kml(tmp_path / 'n.kml', corpo)])
    pastas = [x['pasta'].lower() for x in t]
    assert len(set(pastas)) == len(pastas)


@pytest.mark.parametrize('nome, esperado', [('..', 'talhao'), ('.', 'talhao'),
                                            ('CON', '_CON'), ('aux.txt', '_aux.txt'),
                                            ('Talhão 3.', 'Talhao_3')])
def test_slug_seguro_no_windows(nome, esperado):
    assert C.slug(nome) == esperado


def test_epsg_hemisferio_norte():
    t = [{'anel': [(-61.0, 2.8), (-60.99, 2.8), (-60.99, 2.81), (-61.0, 2.8)]}]
    assert C.epsg_sugerido(t)[0] == 31974          # SIRGAS 2000 / UTM 20N


# ─────────────────────────────────────────────────────── cálculo e validação ──
def test_idw_em_blocos_igual_ao_calculo_direto():
    rng = np.random.default_rng(1)
    t = C.carregar_talhoes([KML])[0]
    pontos, prof, CI, _, _ = C.ler_falker(XLSX)
    idx, _ = C.atribuir_pontos(pontos, [t], 31982)
    sel = idx == 0
    R = C.interpolar(pontos[sel].reset_index(drop=True), CI[:, sel], t['anel'], 31982, 12)
    GX, GY = np.meshgrid(R['gx'], R['gy'])
    k = rng.choice(np.flatnonzero(R['mask']), 20, replace=False)
    D = np.hypot(GX.ravel()[k, None] - R['sx'], GY.ravel()[k, None] - R['sy'])
    W = 1 / np.maximum(D, 1) ** 2
    W /= W.sum(1, keepdims=True)
    esperado = CI[:, sel] @ W.T
    obtido = R['VOL'].reshape(len(prof), -1)[:, k]
    assert np.allclose(obtido, esperado)
    assert np.isnan(R['VOL'][:, ~R['mask']]).all()


def test_validacao_nao_ve_mapa_em_ruido_puro():
    rng = np.random.default_rng(7)
    prof = np.arange(2.5, 20.01, 2.5)
    com_suporte = 0
    for _ in range(200):
        n = 12
        sx, sy = rng.uniform(0, 400, n), rng.uniform(0, 400, n)
        CI = rng.normal(1500, 300, (len(prof), n))
        V = C.validar_camadas(prof, CI, sx, sy, [('0-20', 0, 20)])
        com_suporte += V['0-20']['suporte']
    assert com_suporte / 200 < 0.15


# ────────────────────────────────────────────────────────── projeto completo ──
@pytest.fixture(scope='module')
def projeto(tmp_path_factory):
    pasta = tmp_path_factory.mktemp('p150')
    return C.processar_projeto('Fazenda Exemplo', [KML], XLSX, str(pasta), log=lambda *_: None)


def _resumo(projeto):
    return pd.read_csv(os.path.join(projeto['destino'], 'resumo_do_projeto.csv'),
                       sep=';', decimal=',', encoding='utf-8-sig')


def test_corte_de_profundidade_por_talhao(projeto):
    r = _resumo(projeto).set_index('talhao')
    assert r.loc['Talhão Norte', 'profundidade_max_cm'] == 60      # antes: 52,5
    assert r.loc['Talhão Sul', 'profundidade_max_cm'] == 52.5
    assert 'ci_medio_40-52,5cm' in r.columns                       # rótulo real da camada
    html = open(projeto['htmls'][1], encoding='utf-8').read()
    assert '40 – 52,5 cm' in html and '40 – 60 cm' not in html


def test_csv_abre_no_excel_em_portugues(projeto):
    caminho = os.path.join(projeto['destino'], 'Talhao_Sul', 'perfis_brutos.csv')
    with open(caminho, 'rb') as f:
        cab = f.read(40)
    assert cab.startswith(b'\xef\xbb\xbf') and b';' in cab


def test_talhao_com_erro_nao_derruba_os_outros(tmp_path):
    import shutil
    kml = tmp_path / 't.kml'
    shutil.copy(KML, kml)
    minusculo = ('<Placemark><name>Minúsculo</name><Polygon><outerBoundaryIs><LinearRing>'
                 '<coordinates>-51.30000,-25.30000,0 -51.29998,-25.30000,0 '
                 '-51.29998,-25.29998,0 -51.30000,-25.30000,0</coordinates></LinearRing>'
                 '</outerBoundaryIs></Polygon></Placemark>')
    texto = kml.read_text(encoding='utf-8').replace('</Document>', minusculo + '</Document>')
    kml.write_text(texto, encoding='utf-8')
    # 3 perfis da planilha em cima do talhão minúsculo
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    for c in (COL1 + 30, COL1 + 31, COL1 + 32):
        ws.cell(row=21, column=c).value = -25.299995
        ws.cell(row=22, column=c).value = -51.299987
    x = tmp_path / 'x.xlsx'
    wb.save(x)
    r = C.processar_projeto('P', [str(kml)], str(x), str(tmp_path / 'saida'),
                            log=lambda *_: None)
    assert 'Talhão Norte' in r['talhoes'] and 'Talhão Sul' in r['talhoes']
    assert r['falhas'] == 1
    situacoes = {l['talhao']: l['situacao'] for l in r['resumo']}
    assert situacoes['Minúsculo'].startswith('erro')
    assert os.path.exists(os.path.join(r['destino'], 'resumo_do_projeto.csv'))


def test_texto_do_usuario_nao_vira_codigo(tmp_path):
    import shutil
    kml = tmp_path / 't.kml'
    shutil.copy(KML, kml)
    kml.write_text(kml.read_text(encoding='utf-8').replace(
        'Talhão Norte', 'Lote 7 &lt;A&gt; &lt;/title&gt;&lt;script&gt;x()&lt;/script&gt;'),
        encoding='utf-8')
    x = _planilha(tmp_path, celulas={(13, COL1 + 1): '<img src=x onerror=y()>',
                                     (15, COL1): '<i>14/09'})
    r = C.processar_projeto('Faz. Silva & Filhos <T1>', [str(kml)], x, str(tmp_path / 's'),
                            log=lambda *_: None)
    assert r['falhas'] == 0 and len(r['htmls']) == 2
    html = open(r['htmls'][0], encoding='utf-8').read()
    titulo = html.split('<title>')[1].split('</title>')[0]
    assert '<script>' not in titulo and '&lt;script&gt;' in titulo
    assert '<img src=x' not in html.split('id="payload"')[0]
    for t in r['talhoes']:
        pdf = os.path.join(r['destino'], C.slug(t, 56), C.slug(t, 56) + '.pdf')
        assert os.path.getsize(pdf) > 10000


def test_gpx_valido_com_e_comercial():
    t = C.carregar_talhoes([KML])[0]
    p = planejar(t, 'minimo')
    p['talhao'] = 'Faz. Silva & Filhos <T1>'
    ET.fromstring(_gpx(p))                        # antes: XML inválido


# ───────────────────────────────────────────────── renomear talhões (1.6.0) ──
def test_chave_estavel_e_apelido():
    t = C.carregar_talhoes([KML])
    assert len({x['chave'] for x in t}) == 2
    assert C.carregar_talhoes([KML])[1]['chave'] == t[1]['chave']      # estável
    t2 = C.carregar_talhoes([KML], {t[1]['chave']: 'Gleba do Açude'})
    assert t2[1]['nome'] == 'Gleba do Açude' and t2[1]['nome_kml'] == 'Talhão Sul'
    assert t2[1]['pasta'] == 'Gleba_do_Acude'
    assert t2[0]['nome'] == 'Talhão Norte'


def test_apelido_vale_no_projeto(tmp_path):
    chave = C.carregar_talhoes([KML])[1]['chave']
    r = C.processar_projeto('P', [KML], XLSX, str(tmp_path), log=lambda *_: None,
                            apelidos={chave: 'Gleba do Açude'})
    assert 'Gleba do Açude' in r['talhoes']
    html = open(os.path.join(r['destino'], 'Gleba_do_Acude', 'Gleba_do_Acude.html'),
                encoding='utf-8').read()
    assert 'Gleba do Açude' in html.split('</title>')[0]
    res = {l['talhao']: l for l in r['resumo']}
    assert res['Gleba do Açude']['nome_no_kml'] == 'Talhão Sul'


def test_apelido_vale_na_conferencia():
    chave = C.carregar_talhoes([KML])[0]['chave']
    c = C.conferencia([KML], XLSX, apelidos={chave: 'Lote 1'})
    assert [t['nome'] for t in c['talhoes']] == ['Lote 1', 'Talhão Sul']
